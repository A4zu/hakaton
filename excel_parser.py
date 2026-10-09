"""Excel-dən maliyyə göstəricilərini sətir adına görə tap (EN + AZ).

Düzəlişlər: əvvəl DƏQİQ uyğunluq, sonra qismən; yanlış sətirləri (Cost of sales, Deferred revenue,
Gəlir vergisi...) istisna edir; yalnız etiketdən SAĞDAKI rəqəmləri oxuyur; il başlığı azalan sıradadırsa
(2022, 2021, ...) sıranı çevirir ki, "son rəqəm = ən son il" olsun.
"""
import re
from openpyxl import load_workbook

# sahə: (dəqiq adlar, qismən adlar, istisna sözlər)
KEYS = {
    "current_assets": (["total current assets", "current assets", "cari aktivlər", "dövriyyə aktivləri"],
                       ["total current assets", "current assets", "cari aktivlər", "dövriyyə aktivləri"], []),
    "current_liabilities": (["total current liabilities", "current liabilities", "qısamüddətli öhdəliklər"],
                            ["total current liabilities", "current liabilities", "qısamüddətli öhdəliklər"], []),
    "total_assets": (["total assets", "cəmi aktivlər", "aktivlər cəmi"], ["total assets", "cəmi aktivlər", "aktivlər cəmi"], []),
    "total_liabilities": (["total liabilities", "cəmi öhdəliklər", "öhdəliklər cəmi"],
                          ["total liabilities", "cəmi öhdəliklər", "öhdəliklər cəmi"], ["and equity", "və kapital"]),
    "retained_earnings": (["retained earnings", "bölünməmiş mənfəət"], ["retained earnings", "bölünməmiş mənfəət"], []),
    "ebit": (["ebit", "operating income", "operating profit", "əməliyyat mənfəəti"],
             ["operating income", "operating profit", "əməliyyat mənfəəti"], []),
    "ebitda": (["ebitda"], [], []),
    "sales": (["total revenue", "revenue", "revenues", "net sales", "sales", "satış", "satışlar", "gəlir", "gəlirlər"],
              ["total revenue", "net sales", "satış gəlir"], ["cost", "deferred", "other", "vergi", "tax", "maya", "gözlənilən"]),
    "market_equity": (["market value of equity", "market capitalization", "market cap", "bazar dəyəri"],
                      ["market value of equity", "market capitalization", "bazar dəyəri"], []),
    "debt": (["total debt", "ümumi borc", "cəmi borc"], ["total debt", "ümumi borc", "cəmi borc"], []),
    "annual_payment": (["debt service", "annual debt payment", "illik borc ödənişi"],
                       ["debt service", "annual debt payment", "illik borc ödənişi"], []),
}
REQUIRED = list(KEYS)


def _clean(s):
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _row(row):
    """(etiket, etiketdən sağdakı rəqəmlər)"""
    for i, c in enumerate(row):
        if isinstance(c, str) and c.strip():
            nums = [float(x) for x in row[i + 1:] if isinstance(x, (int, float)) and not isinstance(x, bool)]
            return _clean(c), nums
    return None, []


def _is_year_header(row):
    ys = [c for c in row if isinstance(c, (int, float)) and not isinstance(c, bool) and 1990 <= c <= 2100 and c == int(c)]
    return ys if len(ys) >= 2 and len(ys) >= sum(isinstance(c, (int, float)) for c in row) else None


def _to_num(x):
    """CSV xanası → ədəd (boşluq/minlik ayırıcı, vergül-onluq dəstəklənir); olmazsa özü."""
    if not isinstance(x, str):
        return x
    t = x.strip().replace(" ", "").replace("\u00a0", "")
    if re.fullmatch(r"-?\(?\d[\d.,]*\)?", t):
        neg = t.startswith("(") or t.startswith("-")
        t = t.strip("()-")
        if "," in t and "." in t:
            t = t.replace(",", "")
        elif "," in t:
            t = t.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", t) else t.replace(",", ".")
        try:
            return -float(t) if neg else float(t)
        except ValueError:
            return x
    return x


def _scan(title, rows, found, notes):
    desc = False
    for r in rows:  # il başlığı
        ys = _is_year_header(r)
        if ys:
            desc = ys[0] > ys[-1]
            break
    parsed = [_row(r) for r in rows]
    for mode in ("exact", "partial"):
        for label, nums in parsed:
            if not label or not nums:
                continue
            for field, (exact, partial, excl) in KEYS.items():
                if field in found:
                    continue
                if mode == "exact":
                    hit = label in exact
                else:
                    hit = any(k in label for k in partial) and not any(x in label for x in excl)
                if hit:
                    series = nums[::-1] if desc else nums
                    found[field] = series if field == "ebitda" else series[-1]
    if desc:
        notes.append(f"'{title}': illər azalan sıradadır, çevrildi.")


def parse_excel(file):
    wb = load_workbook(file, data_only=True)
    found, notes = {}, []
    for ws in wb.worksheets:
        _scan(ws.title, list(ws.iter_rows(values_only=True)), found, notes)
    return found, [k for k in REQUIRED if k not in found], notes


def parse_csv(raw: bytes):
    """CSV: ayırıcı (, ; tab) avtomatik tapılır; UTF-8 (BOM ilə də) və cp1251/latin-1 qəbul edilir."""
    import csv, io
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [[_to_num(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    found, notes = {}, []
    _scan("CSV", rows, found, notes)
    return found, [k for k in REQUIRED if k not in found], notes
